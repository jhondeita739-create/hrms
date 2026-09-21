from __future__ import annotations

import math
import textwrap
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets" / "workflows"
OUTPUT = DOCS / "HRMS_Workflows.docx"

NAVY = "0B1F3A"
BLUE = "2457E6"
LIGHT_BLUE = "EAF1FF"
GREEN = "119B6B"
LIGHT_GREEN = "E8F8F1"
AMBER = "D97706"
LIGHT_AMBER = "FFF4DD"
RED = "DC3545"
LIGHT_RED = "FDEBEC"
SLATE = "53657D"
LIGHT_SLATE = "F2F5F9"
WHITE = "FFFFFF"
INK = "142033"


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


class Flowchart:
    def __init__(self, title: str, width: int = 1800, height: int = 1050):
        self.image = Image.new("RGB", (width, height), "#FFFFFF")
        self.draw = ImageDraw.Draw(self.image)
        self.width = width
        self.height = height
        self.nodes: dict[str, tuple[int, int, int, int]] = {}
        self.draw.rectangle((0, 0, width, 92), fill="#0B1F3A")
        self.draw.text((55, 27), title, fill="white", font=font(34, True))

    def label(self, x: int, y: int, text: str, size: int = 22, color: str = INK, bold: bool = False):
        self.draw.text((x, y), text, fill=f"#{color}", font=font(size, bold))

    def node(
        self,
        key: str,
        x: int,
        y: int,
        w: int,
        h: int,
        text: str,
        fill: str = LIGHT_BLUE,
        outline: str = BLUE,
        kind: str = "process",
    ):
        box = (x, y, x + w, y + h)
        self.nodes[key] = box
        if kind == "decision":
            points = [(x + w // 2, y), (x + w, y + h // 2), (x + w // 2, y + h), (x, y + h // 2)]
            self.draw.polygon(points, fill=f"#{fill}", outline=f"#{outline}", width=4)
            text_width = max(13, int(w / 15))
        elif kind == "data":
            self.draw.rounded_rectangle(box, radius=45, fill=f"#{fill}", outline=f"#{outline}", width=4)
            text_width = max(15, int(w / 13))
        else:
            self.draw.rounded_rectangle(box, radius=22, fill=f"#{fill}", outline=f"#{outline}", width=4)
            text_width = max(15, int(w / 13))
        lines = textwrap.wrap(text, width=text_width)
        line_font = font(23, True)
        spacing = 8
        line_boxes = [self.draw.textbbox((0, 0), line, font=line_font) for line in lines]
        total_h = sum(b[3] - b[1] for b in line_boxes) + spacing * max(0, len(lines) - 1)
        cy = y + (h - total_h) / 2
        for line, bounds in zip(lines, line_boxes):
            tw = bounds[2] - bounds[0]
            self.draw.text((x + (w - tw) / 2, cy), line, fill=f"#{INK}", font=line_font)
            cy += bounds[3] - bounds[1] + spacing

    def point(self, key: str, side: str):
        x1, y1, x2, y2 = self.nodes[key]
        return {
            "top": ((x1 + x2) // 2, y1),
            "bottom": ((x1 + x2) // 2, y2),
            "left": (x1, (y1 + y2) // 2),
            "right": (x2, (y1 + y2) // 2),
        }[side]

    def arrow(
        self,
        start: tuple[int, int],
        end: tuple[int, int],
        label: str | None = None,
        via: list[tuple[int, int]] | None = None,
        color: str = SLATE,
    ):
        points = [start, *(via or []), end]
        self.draw.line(points, fill=f"#{color}", width=5, joint="curve")
        previous = points[-2]
        angle = math.atan2(end[1] - previous[1], end[0] - previous[0])
        arrow_size = 16
        p1 = (
            end[0] - arrow_size * math.cos(angle - math.pi / 6),
            end[1] - arrow_size * math.sin(angle - math.pi / 6),
        )
        p2 = (
            end[0] - arrow_size * math.cos(angle + math.pi / 6),
            end[1] - arrow_size * math.sin(angle + math.pi / 6),
        )
        self.draw.polygon([end, p1, p2], fill=f"#{color}")
        if label:
            middle = points[len(points) // 2]
            bounds = self.draw.textbbox((0, 0), label, font=font(19, True))
            tw = bounds[2] - bounds[0]
            self.draw.rounded_rectangle(
                (middle[0] - tw / 2 - 10, middle[1] - 34, middle[0] + tw / 2 + 10, middle[1] - 3),
                radius=8,
                fill="white",
            )
            self.draw.text((middle[0] - tw / 2, middle[1] - 31), label, fill=f"#{color}", font=font(19, True))

    def connect(self, first: str, second: str, first_side: str = "right", second_side: str = "left", **kwargs):
        self.arrow(self.point(first, first_side), self.point(second, second_side), **kwargs)

    def legend(self):
        y = self.height - 58
        items = [(BLUE, "User or HR action"), (GREEN, "Automated success"), (AMBER, "Decision or validation"), (RED, "Terminal or blocked")]
        x = 55
        for color, text in items:
            self.draw.rounded_rectangle((x, y, x + 28, y + 28), radius=6, fill=f"#{color}")
            self.draw.text((x + 40, y + 1), text, fill=f"#{SLATE}", font=font(18))
            x += 300

    def save(self, filename: str):
        self.legend()
        path = ASSETS / filename
        self.image.save(path, quality=96, dpi=(180, 180))
        return path


def end_to_end_chart():
    chart = Flowchart("End-to-end recruitment and employee conversion", height=1100)
    chart.node("vacancy", 55, 165, 230, 105, "Publish vacancy")
    chart.node("apply", 340, 165, 230, 105, "Submit application")
    chart.node("screen", 625, 165, 230, 105, "Resume screening")
    chart.node("hr", 910, 165, 230, 105, "HR interview")
    chart.node("final", 1195, 165, 230, 105, "Final interview")
    chart.node("decision", 1500, 150, 240, 135, "Final decision", LIGHT_AMBER, AMBER, "decision")
    for first, second in [("vacancy", "apply"), ("apply", "screen"), ("screen", "hr"), ("hr", "final"), ("final", "decision")]:
        chart.connect(first, second)

    chart.node("hired", 1420, 410, 250, 100, "Hired", LIGHT_GREEN, GREEN)
    chart.node("rejected", 1050, 410, 250, 100, "Rejected", LIGHT_RED, RED)
    chart.node("withdrawn", 680, 410, 250, 100, "Withdrawn", LIGHT_RED, RED)
    chart.arrow(chart.point("decision", "bottom"), chart.point("hired", "top"), "Hire", via=[(1620, 345), (1545, 345)])
    chart.arrow(chart.point("decision", "left"), chart.point("rejected", "right"), "Reject", via=[(1375, 217), (1375, 460)])
    chart.arrow(chart.point("decision", "left"), chart.point("withdrawn", "right"), "Withdraw", via=[(1345, 320), (980, 320), (980, 460)])

    chart.node("temp", 1420, 610, 250, 100, "Temporary account")
    chart.node("requirements", 1060, 610, 270, 100, "Submit requirements")
    chart.node("training", 700, 610, 270, 100, "Release training")
    chart.node("permanent", 340, 610, 270, 100, "Permanent access", LIGHT_GREEN, GREEN)
    chart.node("onboard", 55, 610, 230, 100, "Employee onboarding", LIGHT_GREEN, GREEN)
    chart.connect("hired", "temp", "bottom", "top")
    chart.connect("temp", "requirements", "left", "right")
    chart.connect("requirements", "training", "left", "right")
    chart.connect("training", "permanent", "left", "right")
    chart.connect("permanent", "onboard", "left", "right")

    chart.label(60, 810, "Key result", 22, BLUE, True)
    chart.node("history", 55, 855, 500, 100, "Every transition writes history and audit data", LIGHT_SLATE, SLATE, "data")
    chart.node("notify", 650, 855, 500, 100, "Applicants receive portal and email updates", LIGHT_SLATE, SLATE, "data")
    chart.node("human", 1245, 855, 500, 100, "Human HR review controls every decision", LIGHT_SLATE, SLATE, "data")
    return chart.save("01_end_to_end.png")


def stage_chart():
    chart = Flowchart("Recruitment stages and transition safeguards", height=1150)
    x = 90
    stages = [
        ("applied", "Applied"),
        ("screen", "Resume Screening"),
        ("hr", "HR Interview"),
        ("final", "Final Interview"),
        ("hired", "Hired"),
    ]
    for index, (key, title) in enumerate(stages):
        fill, border = (LIGHT_GREEN, GREEN) if key == "hired" else (LIGHT_BLUE, BLUE)
        chart.node(key, x + index * 335, 170, 270, 100, title, fill, border)
        if index:
            chart.connect(stages[index - 1][0], key)

    chart.node("verify", 380, 385, 320, 125, "Resume verified?", LIGHT_AMBER, AMBER, "decision")
    chart.node("evaluated", 1045, 385, 330, 125, "Interview completed and evaluated?", LIGHT_AMBER, AMBER, "decision")
    chart.arrow(chart.point("screen", "bottom"), chart.point("verify", "top"))
    chart.arrow(chart.point("hr", "bottom"), chart.point("evaluated", "top"))

    chart.node("block_resume", 120, 610, 360, 100, "Block progression and request verification", LIGHT_RED, RED)
    chart.node("block_interview", 785, 610, 360, 100, "Block hiring and request evaluation", LIGHT_RED, RED)
    chart.arrow(chart.point("verify", "left"), chart.point("block_resume", "top"), "No", via=[(270, 447), (270, 570)])
    chart.arrow(chart.point("evaluated", "left"), chart.point("block_interview", "top"), "No", via=[(930, 447), (930, 570)])

    chart.node("rejected", 1230, 610, 250, 100, "Rejected", LIGHT_RED, RED)
    chart.node("withdrawn", 1510, 610, 250, 100, "Withdrawn", LIGHT_RED, RED)
    chart.label(1230, 760, "Terminal alternatives from any active stage", 22, RED, True)
    chart.node("history", 280, 865, 520, 100, "Stage change records actor, reason, old stage, and new stage", LIGHT_SLATE, SLATE, "data")
    chart.node("special", 1000, 865, 520, 100, "Hired is available only through Hire & onboard", LIGHT_GREEN, GREEN, "data")
    return chart.save("02_recruitment_stages.png")


def application_chart():
    chart = Flowchart("Public application validation and storage", height=1230)
    chart.node("form", 90, 150, 370, 105, "Contact details + PDF resume")
    chart.node("fields", 610, 140, 380, 125, "Fields valid?", LIGHT_AMBER, AMBER, "decision")
    chart.node("pdf", 1140, 140, 420, 125, "Readable resume PDF?", LIGHT_AMBER, AMBER, "decision")
    chart.connect("form", "fields")
    chart.connect("fields", "pdf", label="Yes")
    chart.node("field_error", 440, 380, 370, 100, "Show field-level errors", LIGHT_RED, RED)
    chart.arrow(chart.point("fields", "bottom"), chart.point("field_error", "top"), "No")
    chart.node("pdf_error", 1180, 380, 370, 100, "Explain PDF validation issue", LIGHT_RED, RED)
    chart.arrow(chart.point("pdf", "bottom"), chart.point("pdf_error", "top"), "No")

    chart.node("duplicate", 690, 575, 420, 125, "Duplicate application?", LIGHT_AMBER, AMBER, "decision")
    chart.arrow(chart.point("pdf", "bottom"), chart.point("duplicate", "top"), "Yes", via=[(1350, 535), (900, 535)])
    chart.node("existing", 190, 780, 420, 100, "Return existing tracking reference", LIGHT_SLATE, SLATE)
    chart.node("create", 690, 780, 420, 100, "Create applicant and application", LIGHT_GREEN, GREEN)
    chart.node("storage", 1190, 780, 420, 100, "Store resume and extracted text", LIGHT_GREEN, GREEN)
    chart.arrow(chart.point("duplicate", "left"), chart.point("existing", "top"), "Yes", via=[(480, 637), (480, 740)])
    chart.arrow(chart.point("duplicate", "bottom"), chart.point("create", "top"), "No")
    chart.connect("create", "storage")
    chart.node("reference", 690, 1000, 420, 100, "Display new APL tracking reference", LIGHT_GREEN, GREEN)
    chart.arrow(chart.point("storage", "bottom"), chart.point("reference", "top"), via=[(1400, 950), (900, 950)])
    return chart.save("03_application_validation.png")


def interview_chart():
    chart = Flowchart("Interview scheduling, evaluation, and communication", height=1160)
    chart.label(70, 125, "HR reviewer", 25, BLUE, True)
    chart.label(675, 125, "HRMS + Supabase", 25, GREEN, True)
    chart.label(1300, 125, "Applicant", 25, AMBER, True)
    chart.draw.line((600, 115, 600, 1050), fill="#D6DFEA", width=4)
    chart.draw.line((1220, 115, 1220, 1050), fill="#D6DFEA", width=4)

    chart.node("move", 75, 190, 440, 100, "Move to interview stage")
    chart.node("validate", 680, 190, 440, 100, "Validate verified resume", LIGHT_GREEN, GREEN)
    chart.node("qualified", 1285, 190, 440, 100, "Receive qualification update", LIGHT_AMBER, AMBER)
    chart.connect("move", "validate")
    chart.connect("validate", "qualified")

    chart.node("schedule", 75, 405, 440, 100, "Set date, type, and location")
    chart.node("save", 680, 405, 440, 100, "Save schedule + send email", LIGHT_GREEN, GREEN)
    chart.node("details", 1285, 405, 440, 100, "Receive schedule details", LIGHT_AMBER, AMBER)
    chart.connect("schedule", "save")
    chart.connect("save", "details")

    chart.node("conduct", 75, 620, 440, 100, "Conduct and complete interview")
    chart.node("record", 680, 620, 440, 100, "Store completion status", LIGHT_GREEN, GREEN)
    chart.connect("conduct", "record")

    chart.node("evaluate", 75, 835, 440, 100, "Record recommendation + comments")
    chart.node("decision", 680, 825, 440, 120, "Advance, reject, withdraw, or hire", LIGHT_AMBER, AMBER, "decision")
    chart.connect("evaluate", "decision")
    chart.node("track", 1285, 835, 440, 100, "Track latest application status", LIGHT_SLATE, SLATE)
    chart.connect("decision", "track")
    return chart.save("04_interview_workflow.png")


def hiring_chart():
    chart = Flowchart("Hire and onboard: protected employee conversion", height=1220)
    chart.node("hire", 80, 145, 350, 105, "HR selects Hire & onboard")
    chart.node("gate", 580, 130, 450, 135, "Resume verified and interview evaluated?", LIGHT_AMBER, AMBER, "decision")
    chart.node("blocked", 1210, 145, 420, 105, "Block and explain missing prerequisite", LIGHT_RED, RED)
    chart.connect("hire", "gate")
    chart.connect("gate", "blocked", label="No")

    chart.node("application", 580, 380, 450, 100, "Mark application Hired", LIGHT_GREEN, GREEN)
    chart.arrow(chart.point("gate", "bottom"), chart.point("application", "top"), "Yes")
    chart.node("employee", 120, 585, 420, 100, "Create employee + employment record", LIGHT_GREEN, GREEN)
    chart.node("auth", 690, 570, 420, 125, "Create or link temporary Auth user", LIGHT_GREEN, GREEN)
    chart.node("lifecycle", 1260, 585, 420, 100, "Create preboarding lifecycle", LIGHT_GREEN, GREEN)
    chart.arrow(chart.point("application", "bottom"), chart.point("employee", "top"), via=[(805, 530), (330, 530)])
    chart.connect("employee", "auth")
    chart.connect("auth", "lifecycle")

    chart.node("records", 160, 810, 420, 100, "Create requirements and onboarding")
    chart.node("training", 690, 810, 420, 100, "Prepare locked training schedule")
    chart.node("invite", 1220, 810, 420, 100, "Send secure activation email")
    chart.arrow(chart.point("lifecycle", "bottom"), chart.point("records", "top"), via=[(1470, 760), (370, 760)])
    chart.connect("records", "training")
    chart.connect("training", "invite")
    chart.node("portal", 690, 1020, 420, 100, "Activate password + MFA, then open portal", LIGHT_BLUE, BLUE)
    chart.arrow(chart.point("invite", "bottom"), chart.point("portal", "top"), via=[(1430, 970), (900, 970)])
    return chart.save("05_hire_and_onboard.png")


def promotion_chart():
    chart = Flowchart("Requirements, training release, and permanent access", height=1260)
    chart.node("upload", 90, 145, 360, 105, "Employee uploads requirement")
    chart.node("save", 600, 145, 360, 105, "Store document + version", LIGHT_GREEN, GREEN)
    chart.node("status", 1110, 145, 360, 105, "Set status Submitted", LIGHT_GREEN, GREEN)
    chart.connect("upload", "save")
    chart.connect("save", "status")

    chart.node("all", 170, 390, 410, 130, "All required items satisfied?", LIGHT_AMBER, AMBER, "decision")
    chart.node("ontime", 700, 390, 410, 130, "All submissions on time?", LIGHT_AMBER, AMBER, "decision")
    chart.node("prepared", 1230, 390, 410, 130, "Training prepared?", LIGHT_AMBER, AMBER, "decision")
    chart.arrow(chart.point("status", "bottom"), chart.point("all", "top"), via=[(1290, 330), (375, 330)])
    chart.connect("all", "ontime", label="Yes")
    chart.connect("ontime", "prepared", label="Yes")

    chart.node("temporary", 90, 665, 430, 105, "Keep temporary access and show remaining work", LIGHT_RED, RED)
    chart.arrow(chart.point("all", "left"), chart.point("temporary", "top"), "No", via=[(115, 455), (115, 625)])
    chart.arrow(chart.point("ontime", "bottom"), chart.point("temporary", "right"), "No", via=[(905, 605), (550, 605), (550, 717)])
    chart.arrow(chart.point("prepared", "bottom"), chart.point("temporary", "right"), "No", via=[(1435, 620), (570, 620), (570, 717)])

    chart.node("permanent", 1180, 665, 500, 105, "Promote account to Permanent", LIGHT_GREEN, GREEN)
    chart.arrow(chart.point("prepared", "bottom"), chart.point("permanent", "top"), "Yes")
    chart.node("role", 1180, 860, 500, 100, "Assign employee role + activate profile", LIGHT_GREEN, GREEN)
    chart.connect("permanent", "role", "bottom", "top")
    chart.node("release", 650, 1050, 500, 100, "Release training + mark onboarding ready", LIGHT_GREEN, GREEN)
    chart.arrow(chart.point("role", "bottom"), chart.point("release", "right"), via=[(1430, 1010), (1170, 1010)])
    chart.node("notify", 90, 1050, 450, 100, "Notify employee + write audit log", LIGHT_SLATE, SLATE)
    chart.connect("release", "notify", "left", "right")
    return chart.save("06_access_promotion.png")


def email_chart():
    chart = Flowchart("Email delivery responsibilities", height=1020)
    chart.label(80, 135, "Recruitment notifications", 26, BLUE, True)
    chart.node("event", 80, 200, 390, 100, "Screening, interview, reject, or hire")
    chart.node("portal", 570, 200, 360, 100, "Save portal notification", LIGHT_GREEN, GREEN)
    chart.node("mailer", 1030, 200, 330, 100, "HRMS SMTP mailer", LIGHT_GREEN, GREEN)
    chart.node("gmail", 1460, 200, 260, 100, "Gmail SMTP", LIGHT_AMBER, AMBER)
    chart.connect("event", "portal")
    chart.connect("portal", "mailer")
    chart.connect("mailer", "gmail")

    chart.label(80, 455, "Authentication emails", 26, BLUE, True)
    chart.node("auth_event", 80, 520, 390, 100, "Admin confirmation or employee activation")
    chart.node("supabase", 650, 520, 390, 100, "Supabase Auth creates secure link", LIGHT_GREEN, GREEN)
    chart.node("custom", 1220, 520, 390, 100, "Supabase custom Gmail SMTP", LIGHT_AMBER, AMBER)
    chart.connect("auth_event", "supabase")
    chart.connect("supabase", "custom")

    chart.node("inbox", 680, 755, 440, 110, "Recipient inbox", LIGHT_GREEN, GREEN)
    chart.arrow(chart.point("gmail", "bottom"), chart.point("inbox", "right"), via=[(1590, 700), (1150, 700), (1150, 810)])
    chart.arrow(chart.point("custom", "bottom"), chart.point("inbox", "right"), via=[(1415, 720), (1170, 720), (1170, 810)])
    return chart.save("07_email_delivery.png")


def set_cell_shading(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def add_bullets(document: Document, items: list[str]):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_chart(document: Document, image_path: Path, caption: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(7.15))
    caption_p = document.add_paragraph(caption)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.style = document.styles["Caption"]


def add_rules_table(document: Document, rows: list[tuple[str, str]]):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, title in enumerate(("Workflow point", "Implemented rule")):
        header.cells[index].text = title
        set_cell_shading(header.cells[index], NAVY)
        for run in header.cells[index].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    table.columns[0].width = Inches(2.05)
    table.columns[1].width = Inches(5.0)
    document.add_paragraph()


def build_document(charts: list[Path]):
    document = Document()
    props = document.core_properties
    props.title = "HRMS Recruitment and Onboarding Workflows"
    props.subject = "Functional workflow and process flowcharts"
    props.author = "Priority Handling Logistics, Inc."
    props.keywords = "HRMS, recruitment, onboarding, workflow, flowchart"

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Title", 34, NAVY),
        ("Heading 1", 22, NAVY),
        ("Heading 2", 16, BLUE),
        ("Heading 3", 12, SLATE),
    ]:
        styles[style_name].font.name = "Arial"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.color.rgb = RGBColor.from_string(color)
        styles[style_name].font.bold = True

    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Recruitment & Onboarding\nWorkflows")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Priority Handling Logistics, Inc. • HRMS")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.bold = True
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Implementation reference • Updated {date.today().isoformat()}")
    document.add_paragraph()
    summary = document.add_table(rows=1, cols=1)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.style = "Table Grid"
    cell = summary.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    cell.text = (
        "Purpose: provide a single, implementation-aligned reference for applicant intake, "
        "recruitment decisions, interviews, notifications, hiring, temporary employee access, "
        "requirements, training, permanent access, and MFA."
    )
    document.add_page_break()

    document.add_heading("Contents", level=1)
    add_numbered(document, [
        "End-to-end recruitment overview",
        "Public application and resume validation",
        "Recruitment stages and safeguards",
        "Resume review and AI-assisted ranking",
        "Interview workflow",
        "Applicant notifications and tracking",
        "Hire and temporary-account creation",
        "Requirements and permanent-access promotion",
        "Authentication, MFA, and email responsibilities",
        "Route ownership and acceptance testing",
    ])
    document.add_paragraph(
        "Diagram colors: blue represents a user or HR action, green an automated success, "
        "amber a decision or validation, and red a blocked or terminal result."
    )
    document.add_page_break()

    document.add_heading("1. End-to-end recruitment overview", level=1)
    document.add_paragraph(
        "The workflow begins with a published vacancy and ends when a successful candidate "
        "has permanent employee access and enters onboarding. Rejected and withdrawn "
        "applications retain their history for authorized review."
    )
    add_chart(document, charts[0], "Figure 1. Complete recruitment-to-employee lifecycle")

    document.add_heading("2. Public application and resume validation", level=1)
    document.add_paragraph(
        "The public form intentionally collects only the information required for initial "
        "screening: full name, phone, email, location, privacy consent, and a PDF resume."
    )
    add_chart(document, charts[2], "Figure 2. Application validation and data creation")
    add_rules_table(document, [
        ("Phone", "Accept 7 to 15 digits."),
        ("Resume format", "Accept PDF only, up to 4 MB."),
        ("Resume integrity", "Require a readable, searchable, unencrypted PDF with recognizable resume content."),
        ("Duplicates", "Prevent the same applicant email from applying twice to one vacancy."),
        ("Tracking", "Return an APL reference after a successful submission."),
    ])

    document.add_heading("3. Recruitment stages and safeguards", level=1)
    add_chart(document, charts[1], "Figure 3. Active stages and protected transitions")
    add_rules_table(document, [
        ("Applied", "Assigned automatically after a successful application."),
        ("Resume Screening", "HR verifies the submitted resume before further progression."),
        ("HR Interview", "Scheduling is available only after resume verification."),
        ("Final Interview", "HR records the final interview and evaluation."),
        ("Hired", "Available only through Hire & onboard, not the ordinary stage selector."),
        ("Rejected / Withdrawn", "Terminal outcomes with a recorded reason and stage history."),
    ])

    document.add_heading("4. Resume review and AI-assisted ranking", level=1)
    document.add_paragraph(
        "HR may generate an explainable score from the vacancy criteria and extracted resume "
        "text. The output includes a 0–100 score, recommendation, strengths, concerns, and a "
        "summary. It is decision support only."
    )
    add_bullets(document, [
        "Every applicant remains visible regardless of score.",
        "The system does not automatically reject or hire an applicant.",
        "An authorized reviewer records whether the guidance was accepted or overridden.",
        "Human review controls every recruitment-stage transition.",
    ])

    document.add_heading("5. Interview workflow", level=1)
    add_chart(document, charts[3], "Figure 4. Interview scheduling, communication, and evaluation")
    add_bullets(document, [
        "HR records interview type, start and end time, timezone, location, or meeting URL.",
        "The applicant receives qualification and scheduling updates when SMTP is available.",
        "HR can edit or cancel an interview before completion.",
        "A completed interview and saved evaluation are required before hiring.",
    ])

    document.add_heading("6. Applicant notifications and tracking", level=1)
    document.add_paragraph(
        "Every supported recruitment event creates a portal notification. When Gmail SMTP is "
        "configured, the HRMS also creates an email-delivery record, sends the message, and "
        "stores the provider message ID or delivery error."
    )
    add_bullets(document, [
        "Screening passed",
        "Qualified for interview",
        "Interview scheduled",
        "Rejected",
        "Hired",
    ])
    document.add_paragraph(
        "Applicants use /careers/track with the application reference and matching email to "
        "view the current stage and updates. Tracking remains available if external email fails."
    )

    document.add_heading("7. Hire and temporary-account creation", level=1)
    add_chart(document, charts[4], "Figure 5. Protected Hire & onboard transaction")
    add_rules_table(document, [
        ("Hiring prerequisites", "Verified resume and a completed, evaluated interview."),
        ("Employee record", "Create the employee and current position/department assignment."),
        ("Auth account", "Create or link an employee-compatible Supabase Auth user."),
        ("Temporary role", "Assign preboarding access, requirements, onboarding, and locked training."),
        ("Activation", "Send a secure Supabase invitation through configured custom SMTP."),
    ])

    document.add_heading("8. Requirements and permanent-access promotion", level=1)
    add_chart(document, charts[5], "Figure 6. Automatic temporary-to-permanent access promotion")
    document.add_heading("Requirement states", level=2)
    document.add_paragraph(
        "Pending → Submitted → Under review → Verified. HR may reject a submitted file, after "
        "which the employee uploads a replacement. HR may waive a requirement when no file is needed."
    )
    add_bullets(document, [
        "Submitted, Under review, Verified, and Rejected states require an attached document.",
        "Replacement uploads create document versions.",
        "HR sees submissions in Preboarding and Employee records.",
        "The employee portal hides completed submissions while preserving HR records.",
    ])
    document.add_heading("Promotion gate", level=2)
    document.add_paragraph(
        "Permanent access is granted when all required items were submitted on time and are not "
        "Pending or Rejected, or were waived, and at least one prepared training schedule exists. "
        "The system then assigns the employee role, activates the profile, releases training, "
        "marks onboarding ready, notifies the employee, and writes an audit record."
    )

    document.add_heading("9. Authentication, MFA, and email responsibilities", level=1)
    add_chart(document, charts[6], "Figure 7. Separate application and Supabase Auth email channels")
    add_rules_table(document, [
        ("Applicant updates", "HRMS server sends through SMTP_* Gmail settings."),
        ("Admin confirmation", "Supabase Auth sends through its custom Gmail SMTP settings."),
        ("Employee activation", "Supabase Auth sends the invitation or activation link."),
        ("MFA", "Protected HR and employee sessions proceed through enrollment or verification to AAL2."),
        ("Admin registration", "Public first-administrator registration is disabled."),
    ])

    document.add_heading("10. Route ownership and acceptance testing", level=1)
    add_rules_table(document, [
        ("/careers", "Browse published vacancies."),
        ("/careers/[vacancyId]", "Submit a public application."),
        ("/careers/track", "Track an application by reference and email."),
        ("/hr/recruitment/applicants", "Review applicants and current stages."),
        ("Applicant detail", "Verify resume, assess match, manage interviews, and decide outcome."),
        ("/hr/preboarding", "Manage temporary accounts, requirements, training, and access."),
        ("/employee/onboarding", "Submit employee documents and view unlocked training."),
        ("/hr/records", "Review retained employee records and submitted files."),
    ])
    document.add_heading("Recommended end-to-end acceptance test", level=2)
    add_numbered(document, [
        "Publish a test vacancy and submit a valid searchable PDF resume.",
        "Use the returned reference in Track application.",
        "Verify the resume and generate the advisory match assessment.",
        "Advance to HR Interview and confirm the notification record.",
        "Schedule, complete, and evaluate the interview.",
        "Advance to Final Interview, then use Hire & onboard.",
        "Confirm employee, employment, Auth, lifecycle, requirement, training, and audit records.",
        "Activate the temporary account, complete MFA, and upload every requirement.",
        "Confirm HR can download and review each submission.",
        "Confirm on-time completion releases training and grants Permanent access.",
    ])

    document.add_paragraph()
    closing = document.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("End of workflow reference")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(SLATE)

    document.save(OUTPUT)


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    charts = [
        end_to_end_chart(),
        stage_chart(),
        application_chart(),
        interview_chart(),
        hiring_chart(),
        promotion_chart(),
        email_chart(),
    ]
    build_document(charts)
    print(OUTPUT)


if __name__ == "__main__":
    main()
